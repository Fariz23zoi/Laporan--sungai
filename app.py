import streamlit as st
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, Border, Side
from datetime import datetime
from io import BytesIO
from PIL import Image as PILImage
from concurrent.futures import ThreadPoolExecutor
from functools import lru_cache
import tempfile
import math
import re

st.title("📋 Laporan Penelusuran Sungai PRO MAX")

# ======================
# BULAN
# ======================
bulan_list = [
    "Januari","Februari","Maret","April","Mei","Juni",
    "Juli","Agustus","September","Oktober","November","Desember"
]

bulan = st.selectbox("Bulan", bulan_list)
tahun = st.number_input("Tahun", value=2026)

bulan_romawi = {
    "Januari": "I","Februari": "II","Maret": "III","April": "IV",
    "Mei": "V","Juni": "VI","Juli": "VII","Agustus": "VIII",
    "September": "IX","Oktober": "X","November": "XI","Desember": "XII"
}

# ======================
# SESSION
# ======================
if "data" not in st.session_state:
    st.session_state.data = []

st.subheader("Input Data")

uraian = st.text_input("Uraian")
lokasi = st.text_area("Lokasi")
koordinat = st.text_input("Koordinat")
keterangan = st.text_area("Keterangan")
foto = st.file_uploader("Upload Foto", type=["jpg","png","jpeg"])

col1, col2 = st.columns(2)

with col1:
    if st.button("➕ Tambah Data"):
        if not uraian or not lokasi:
            st.warning("Uraian & Lokasi wajib!")
        elif not foto:
            st.warning("Upload foto dulu!")
        else:
            st.session_state.data.append({
                "uraian": uraian,
                "lokasi": lokasi,
                "koordinat": koordinat,
                "keterangan": keterangan,
                "foto": foto
            })
            st.success("Data ditambahkan!")

with col2:
    if st.button("🗑 Reset Data"):
        st.session_state.data = []

st.write(f"Jumlah data: {len(st.session_state.data)}")

# Preview
if st.session_state.data:
    st.dataframe([
        {k:v for k,v in d.items() if k!="foto"}
        for d in st.session_state.data
    ])

# ======================
# COMPRESS + CACHE
# ======================
@lru_cache(maxsize=200)
def compress_image_cached(file_bytes, quality=65):
    img = PILImage.open(BytesIO(file_bytes))

    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")

    img.thumbnail((1280, 720))

    buffer = BytesIO()
    img.save(buffer, format="JPEG", quality=quality, optimize=True)
    return buffer.getvalue()

# ======================
# EXPORT
# ======================
if st.button("📊 Export Excel SUPER CEPAT"):

    data = st.session_state.data
    total_data = len(data)

    if total_data == 0:
        st.warning("Belum ada data!")
        st.stop()

    # ======================
    # PARALLEL COMPRESS
    # ======================
    def process_image(d):
        d["foto"].seek(0)
        return compress_image_cached(d["foto"].read())

    with ThreadPoolExecutor(max_workers=4) as executor:
        compressed_images = list(executor.map(process_image, data))

    # ======================
    # EXCEL
    # ======================
    wb = Workbook()
    ws = wb.active

    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_top = Alignment(horizontal='left', vertical='top', wrap_text=True)
    bold = Font(bold=True)

    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # KOLOM
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 22
    ws.column_dimensions['E'].width = 40
    ws.column_dimensions['F'].width = 30

    # PAGE SETUP (BIAR PRINT AMAN)
    ws.page_setup.fitToWidth = 1

    # PAGINATION
    data_per_page = 5
    total_halaman = math.ceil(total_data / data_per_page)

    # NAMA SUNGAI
    uraian_text = data[0]["uraian"]
    match = re.search(r"sungai\s+(.+)", uraian_text.lower())
    nama_sungai = match.group(1).upper() if match else uraian_text.upper()

    row_global = 1

    for page in range(total_halaman):

        if page > 0:
            ws.page_breaks.append(row_global)

        # JUDUL
        ws.merge_cells(start_row=row_global, start_column=1, end_row=row_global, end_column=6)
        ws.cell(row=row_global, column=1, value=f"LAPORAN PENELUSURAN SUNGAI {nama_sungai}").alignment = center
        ws.cell(row=row_global, column=1).font = bold

        ws.merge_cells(start_row=row_global+1, start_column=1, end_row=row_global+1, end_column=6)
        ws.cell(row=row_global+1, column=1, value="KOTA/KAB. CIAMIS").alignment = center

        ws.merge_cells(start_row=row_global+2, start_column=1, end_row=row_global+2, end_column=6)
        ws.cell(row=row_global+2, column=1, value="OPERASI DAN PEMELIHARAAN SDA III").alignment = center

        ws.merge_cells(start_row=row_global+3, start_column=1, end_row=row_global+3, end_column=6)
        ws.cell(row=row_global+3, column=1, value=f"BULAN {bulan.upper()} {tahun}").alignment = center

        nomor = f"{page+1}/{total_halaman}/OPSDA-03/{bulan_romawi[bulan]}/{tahun}"
        ws.merge_cells(start_row=row_global+4, start_column=1, end_row=row_global+4, end_column=6)
        ws.cell(row=row_global+4, column=1, value=f"Nomor: {nomor}").alignment = center

        # HEADER
        headers = ["NO","URAIAN","LOKASI","KOORDINAT","FOTO","KETERANGAN"]
        for col, val in enumerate(headers, 1):
            c = ws.cell(row=row_global+5, column=col)
            c.value = val
            c.font = bold
            c.alignment = center
            c.border = border

        # DATA PER PAGE
        start = page * data_per_page
        end = start + data_per_page
        page_data = data[start:end]

        row = row_global + 7

        for i, d in enumerate(page_data, start=1):

            ws.row_dimensions[row].height = 170

            ws.cell(row=row, column=1, value=i)
            ws.cell(row=row, column=2, value=d["uraian"])
            ws.cell(row=row, column=3, value=d["lokasi"])
            ws.cell(row=row, column=4, value=d["koordinat"])
            ws.cell(row=row, column=6, value=d["keterangan"])

            for col in [1,2,3,4,6]:
                ws.cell(row=row, column=col).alignment = left_top
                ws.cell(row=row, column=col).border = border

            # FOTO
            img_bytes = compressed_images[start + i - 1]

            with tempfile.NamedTemporaryFile(delete=False, suffix=".jpg") as tmp:
                tmp.write(img_bytes)
                img = XLImage(tmp.name)

            img.width = 260
            img.height = 150
            img.anchor = f"E{row}"
            ws.add_image(img)

            ws.cell(row=row, column=5).border = border

            row += 1

        # TTD
        ttd_row = row + 2
        today = datetime.now()
        tanggal = f"Ciamis, {today.day} {bulan_list[today.month-1]} {today.year}"

        ws.merge_cells(start_row=ttd_row, start_column=5, end_row=ttd_row, end_column=6)
        ws.cell(row=ttd_row, column=5, value=tanggal).alignment = center

        ws.merge_cells(start_row=ttd_row+1, start_column=5, end_row=ttd_row+1, end_column=6)
        ws.cell(row=ttd_row+1, column=5, value="Juru Sungai OPSDA 03").alignment = center

        ws.merge_cells(start_row=ttd_row+4, start_column=5, end_row=ttd_row+4, end_column=6)
        ws.cell(row=ttd_row+4, column=5, value="Fariz Rionaldi").alignment = center

        row_global = ttd_row + 6

    # ======================
    # SAVE MEMORY (RINGAN)
    # ======================
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Laporan_Sungai_{bulan}_{tahun}.xlsx"

    st.download_button(
        "⬇️ Download Excel",
        data=output,
        file_name=filename
        )
